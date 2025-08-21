# -*- coding: utf-8 -*-
"""
올리브영 '브랜드 랭킹' 자동 수집 (GitHub Actions 전용)
우선순위:
  1) Oxylabs Realtime API(Universal)로 렌더된 HTML 수신
  2) Web Unblocker 프록시 + Playwright로 페이지 로드 & 네트워크 JSON 파싱
  3) 둘 다 실패 시 디버그 아티팩트 남기고 슬랙에 실패 알림

결과:
  - data/올리브영_브랜드_순위.xlsx (월 시트 자동 생성/갱신)
  - 슬랙 Top10 (전일 대비 등락 표기)
  - debug: brand_debug.html/png, brand_api_response.json
"""

import os, re, json, asyncio
from collections import OrderedDict
from datetime import datetime, timedelta, date
from zoneinfo import ZoneInfo
from calendar import monthrange

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font

# ----- Playwright (프록시 폴백용)
from urllib.parse import urlparse
from playwright.async_api import async_playwright

# ----- Google Drive (선택)
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload
from google.oauth2.credentials import Credentials


# =========================
# 설정/시크릿
# =========================
KST = ZoneInfo("Asia/Seoul")
URL = "https://m.oliveyoung.co.kr/m/mtn?menu=ranking&tab=brands"

OUTPUT_DIR = "data"
XLSX_NAME  = "올리브영_브랜드_순위.xlsx"
OUTPUT_PATH = os.path.join(OUTPUT_DIR, XLSX_NAME)

# env
SCRAPING_API = os.environ.get("SCRAPING_API", "").lower()  # "oxylabs" 권장
OXY_USER     = os.environ.get("OXY_USER", "")
OXY_PASS     = os.environ.get("OXY_PASS", "")

PROXY_SERVER = os.environ.get("PROXY_SERVER", "")

SLACK_WEBHOOK_URL    = os.environ.get("SLACK_WEBHOOK_URL", "")
GDRIVE_FOLDER_ID     = os.environ.get("GDRIVE_FOLDER_ID", "")
GOOGLE_CLIENT_ID     = os.environ.get("GOOGLE_CLIENT_ID", "")
GOOGLE_CLIENT_SECRET = os.environ.get("GOOGLE_CLIENT_SECRET", "")
GOOGLE_REFRESH_TOKEN = os.environ.get("GOOGLE_REFRESH_TOKEN", "")

os.makedirs(OUTPUT_DIR, exist_ok=True)


# =========================
# 공통 유틸/정규화
# =========================
CODE_PATTERNS = [re.compile(r"^[A-Z]\d{4,}$"), re.compile(r"^\d{4,}$")]

def normalize_brand_text(t: str) -> str | None:
    if not isinstance(t, str) or not t:
        return None
    s = re.sub(r"[\u200b\ufeff]", "", t)
    s = re.sub(r"\s+", " ", s).strip()
    # 불필요 꼬리표 컷
    s = re.sub(r"\s*(브랜드\s*썸네일|로고.*|이미지.*|타이틀.*)$", "", s).strip()
    # 코드/숫자/과도한 길이 컷
    for p in CODE_PATTERNS:
        if p.match(s):
            return None
    if re.search(r"\d", s):
        return None
    if len(s) > 30 or len(s.split()) > 6:
        return None
    if len(s) == 1 and not re.fullmatch(r"[가-힣]", s):
        return None
    if s.lower() in {"brand","logo","image","title","브랜드","이미지","타이틀"}:
        return None
    return s

def unique_keep_order(xs):
    return list(OrderedDict.fromkeys(xs))

# =========================
# 1) Realtime API (우선)
# =========================
def fetch_html_via_oxylabs(url: str) -> str | None:
    if SCRAPING_API != "oxylabs" or not (OXY_USER and OXY_PASS):
        return None
    try:
        payload = {
            "source": "universal",
            "url": url,
            "render": "html",
            "geo_location": "South Korea",
            "user_agent_type": "mobile",
        }
        r = requests.post(
            "https://realtime.oxylabs.io/v1/queries",
            auth=(OXY_USER, OXY_PASS),  # Basic Auth
            headers={"Content-Type": "application/json"},
            data=json.dumps(payload),
            timeout=60,
        )
        # 디버그 저장(에러 바디 포함)
        with open(os.path.join(OUTPUT_DIR, "brand_api_response.json"), "w", encoding="utf-8") as f:
            try:
                f.write(json.dumps(r.json(), ensure_ascii=False, indent=2))
            except Exception:
                f.write(r.text[:200000])

        r.raise_for_status()
        data = r.json()
        results = data.get("results") or []
        if results and "content" in results[0]:
            return results[0]["content"]
        print("[Oxylabs] content 없음")
    except Exception as e:
        print(f"[Oxylabs] 요청 실패: {e}")
    return None

def extract_brands_from_html(html: str) -> list[str]:
    if not html:
        return []
    names = []
    for m in re.finditer(r'brandsInfo"\s*:\s*{[^}]*"brandName"\s*:\s*"([^"]+)"', html):
        nm = normalize_brand_text(m.group(1))
        if nm:
            names.append(nm)
    if not names:
        for m in re.finditer(r'"brandName"\s*:\s*"([^"]+)"', html):
            nm = normalize_brand_text(m.group(1))
            if nm:
                names.append(nm)
    return unique_keep_order(names)[:100]

# =========================
# 2) Web Unblocker 프록시 + Playwright (폴백)
# =========================
def parse_proxy(proxy_url: str) -> dict | None:
    if not proxy_url:
        return None
    u = urlparse(proxy_url)
    if not (u.scheme and u.hostname and u.port):
        return None
    out = {"server": f"{u.scheme}://{u.hostname}:{u.port}"}
    if u.username:
        out["username"] = u.username
    if u.password:
        out["password"] = u.password
    return out

def is_cf_block_html(html: str) -> bool:
    return bool(html) and (("사람인지 확인" in html) or ("Cloudflare" in html and "확인" in html))

async def fetch_brands_via_playwright(url: str, proxy_url: str) -> list[str]:
    proxy = parse_proxy(proxy_url)
    json_payloads = []
    async with async_playwright() as p:
        iphone = p.devices.get("iPhone 13 Pro")
        browser = await p.chromium.launch(headless=True, proxy=proxy)
        context = await browser.new_context(
            ignore_https_errors=True,   # 프록시 체인 TLS 에러 무시
            **iphone,
            locale="ko-KR",
            extra_http_headers={
                "X-Oxylabs-Geo-Location": "South Korea",
                "X-Oxylabs-Render": "html",
                "X-Oxylabs-Device-Type": "mobile",
            },
        )
        page = await context.new_page()

        async def on_response(resp):
            try:
                ctype = resp.headers.get("content-type","").lower()
                if "application/json" in ctype:
                    url_l = resp.url.lower()
                    if any(k in url_l for k in ["brand","ranking","best"]):
                        jo = await resp.json()
                        json_payloads.append(jo)
            except Exception:
                pass
        page.on("response", lambda r: asyncio.create_task(on_response(r)))

        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=60_000)
        except Exception as e:
            print(f"[Playwright] goto 실패: {e}")

        # 디버그 HTML/PNG 항상 저장
        try:
            html = await page.content()
            with open(os.path.join(OUTPUT_DIR, "brand_debug.html"), "w", encoding="utf-8") as f:
                f.write(html[:200000])
            await page.screenshot(path=os.path.join(OUTPUT_DIR, "brand_debug.png"), full_page=True)
        except Exception:
            pass

        # 간단 스크롤+더보기 (명시적 버튼 클릭 없이 로딩 유도)
        try:
            for _ in range(10):
                await page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                await page.wait_for_timeout(800)
        except Exception:
            pass

        # JSON에서 추출
        names = extract_brands_from_json(json_payloads)

        # 부족하면 HTML 백업 파싱
        if len(names) < 50:
            try:
                html2 = await page.content()
                names = unique_keep_order(names + extract_brands_from_html(html2))[:100]
            except Exception:
                pass

        await context.close()
        await browser.close()
        return names

def extract_brands_from_json(json_objs) -> list[str]:
    names = []
    def walk(node):
        if isinstance(node, dict):
            bi = node.get("brandsInfo") or node.get("brandInfo")
            if isinstance(bi, dict):
                nm = (
                    bi.get("brandName") or bi.get("brandNm")
                    or bi.get("brandKrName") or bi.get("brand_kor_name")
                )
                nm = normalize_brand_text(nm)
                if nm:
                    names.append(nm)
            for v in node.values():
                if isinstance(v, (dict, list)):
                    walk(v)
        elif isinstance(node, list):
            for it in node:
                walk(it)
    for jo in json_objs:
        try:
            walk(jo)
        except Exception:
            pass
    return unique_keep_order(names)[:100]

# =========================
# 엑셀 (월 시트 자동)
# =========================
def month_sheet_name(dt: datetime) -> str:
    return f"{dt.strftime('%y')}년 {dt.month}월"

def ensure_month_sheet(wb, dt: datetime):
    name = month_sheet_name(dt)
    if name in wb.sheetnames:
        ws = wb[name]
    else:
        ws = wb.create_sheet(title=name)
        setup_layout(ws, dt)
    return ws

def setup_layout(ws, dt: datetime):
    last_day = monthrange(dt.year, dt.month)[1]
    ws["A1"] = "브랜드 순위 (올리브영 앱 기준)"
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + last_day)

    ws["A2"] = "일자"
    for d in range(1, last_day + 1):
        ws.cell(row=2, column=1 + d).value = f"{d}일"

    ws["A3"] = "요일"
    for d in range(1, last_day + 1):
        wd = date(dt.year, dt.month, d).weekday()
        ws.cell(row=3, column=1 + d).value = ["월","화","수","목","금","토","일"][wd]

    ws["A4"] = "비고"
    for r in range(1, 101):
        ws.cell(row=4 + r, column=1).value = r

    for r in range(1, 5 + 100):
        for c in range(1, 1 + last_day + 1):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 8
    for d in range(1, last_day + 1):
        ws.column_dimensions[get_column_letter(1 + d)].width = 18

def write_today(ws, now: datetime, brands):
    col = 1 + now.day
    for i in range(100):
        ws.cell(row=5 + i, column=col).value = brands[i] if i < len(brands) else None

def read_rank_map(ws, day: int):
    col = 1 + day
    ranks = {}
    for i in range(100):
        name = ws.cell(row=5 + i, column=col).value
        if name:
            ranks[str(name).strip()] = i + 1
    return ranks

def get_yesterday_rank_map(wb, now: datetime):
    y = now - timedelta(days=1)
    sn = month_sheet_name(y)
    if sn in wb.sheetnames:
        try:
            return read_rank_map(wb[sn], y.day)
        except Exception:
            pass
    return {}

def save_excel_and_get_yesterday_map(brands):
    now = datetime.now(KST)
    if os.path.exists(OUTPUT_PATH):
        wb = load_workbook(OUTPUT_PATH)
    else:
        wb = Workbook()
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
            wb.remove(wb["Sheet"])
    ws = ensure_month_sheet(wb, now)
    if ws.cell(row=2, column=2).value is None:
        setup_layout(ws, now)
    ymap = get_yesterday_rank_map(wb, now)
    write_today(ws, now, brands)
    wb.save(OUTPUT_PATH)
    return ymap, now

# =========================
# Slack
# =========================
def build_delta(today_rank, yesterday_rank):
    if yesterday_rank is None:
        return "(new)"
    diff = yesterday_rank - today_rank
    return f"(↑{diff})" if diff > 0 else (f"(↓{abs(diff)})" if diff < 0 else "(-)")

def post_slack_top10(brands, ymap, now):
    if not SLACK_WEBHOOK_URL:
        return
    if not brands:
        payload = {"text": f"❗ 수집 실패 — {now.strftime('%Y-%m-%d')} (KST) / 디버그 아티팩트를 확인하세요."}
        try:
            requests.post(SLACK_WEBHOOK_URL, data=json.dumps(payload),
                          headers={"Content-Type":"application/json"}, timeout=10)
        except Exception:
            pass
        return

    top10 = brands[:10]
    lines = []
    for idx, nm in enumerate(top10, start=1):
        delta = build_delta(idx, ymap.get(nm))
        lines.append(f"{idx}. {delta} {nm}")
    title = f"📊 올리브영 데일리 브랜드 랭킹 Top10 — {now.strftime('%Y-%m-%d')} (KST)"
    payload = {
        "blocks": [
            {"type": "header", "text": {"type":"plain_text","text":title,"emoji":True}},
            {"type": "section", "text": {"type":"mrkdwn","text":"\n".join(lines)}},
        ]
    }
    try:
        requests.post(SLACK_WEBHOOK_URL, data=json.dumps(payload),
                      headers={"Content-Type":"application/json"}, timeout=10)
    except Exception:
        pass

# =========================
# Google Drive 업로드(선택)
# =========================
def build_drive_service():
    if not (GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET and GOOGLE_REFRESH_TOKEN and GDRIVE_FOLDER_ID):
        return None
    try:
        creds = Credentials(
            token=None,
            refresh_token=GOOGLE_REFRESH_TOKEN,
            token_uri="https://oauth2.googleapis.com/token",
            client_id=GOOGLE_CLIENT_ID,
            client_secret=GOOGLE_CLIENT_SECRET,
            scopes=["https://www.googleapis.com/auth/drive.file"],
        )
        return build("drive", "v3", credentials=creds)
    except Exception as e:
        print(f"[드라이브] 서비스 생성 실패: {e}")
        return None

def find_file_in_folder(service, folder_id, name):
    q = f"name = '{name}' and '{folder_id}' in parents and trashed=false"
    res = service.files().list(q=q, fields="files(id,name)", pageSize=1).execute()
    files = res.get("files", [])
    return files[0]["id"] if files else None

def upload_or_update_to_drive(filepath, folder_id):
    service = build_drive_service()
    if not service:
        return
    try:
        file_id = find_file_in_folder(service, folder_id, os.path.basename(filepath))
        media = MediaFileUpload(
            filepath,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resumable=True,
        )
        if file_id:
            service.files().update(fileId=file_id, media_body=media).execute()
            print(f"[드라이브] 기존 파일 갱신 완료")
        else:
            meta = {"name": os.path.basename(filepath), "parents": [folder_id]}
            service.files().create(body=meta, media_body=media, fields="id").execute()
            print(f"[드라이브] 새 파일 업로드 완료")
    except Exception as e:
        print(f"[드라이브] 업로드 실패: {e}")

# =========================
# main
# =========================
async def run():
    brands = []

    # 1) Realtime API (권장 경로)
    html = fetch_html_via_oxylabs(URL)
    if html:
        with open(os.path.join(OUTPUT_DIR, "brand_debug.html"), "w", encoding="utf-8") as f:
            f.write(html[:200000])
        brands = extract_brands_from_html(html)

    # 2) 폴백: Web Unblocker 프록시 + Playwright
    if not brands and PROXY_SERVER:
        print("[info] Realtime 실패 또는 미설정 — 프록시 폴백 시도")
        try:
            brands = await fetch_brands_via_playwright(URL, PROXY_SERVER)
        except Exception as e:
            print(f"[Playwright] 폴백 실패: {e}")

    # 결과 처리
    now = datetime.now(KST)
    ymap, now = save_excel_and_get_yesterday_map(brands)
    post_slack_top10(brands, ymap, now)
    upload_or_update_to_drive(OUTPUT_PATH, GDRIVE_FOLDER_ID)

if __name__ == "__main__":
    asyncio.run(run())
